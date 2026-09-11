# -*- coding: utf-8 -*-
"""
文档解析器：txt / md / pdf / docx → 纯文本。
每种格式独立函数，解析异常向上抛出让状态机标记 failed。
"""
import os

from docx import Document as DocxDocument
from pypdf import PdfReader


def parse_document(file_path: str, ext: str) -> str:
    """按扩展名解析文档为纯文本"""
    ext = ext.lower()
    if ext in (".txt", ".md"):
        return _parse_text(file_path)
    if ext == ".pdf":
        return _parse_pdf(file_path)
    if ext == ".docx":
        return _parse_docx(file_path)
    raise ValueError(f"不支持的文档格式: {ext}")


def _parse_text(file_path: str) -> str:
    """txt/md：优先 utf-8，回退 gbk"""
    for encoding in ("utf-8", "gbk", "utf-16"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本文件编码")


def _parse_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    text = "\n".join(pages)
    if not text.strip():
        raise ValueError("PDF 未提取到文本（可能为扫描件，暂不支持 OCR）")
    return text


def _parse_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # 表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ValueError("docx 未提取到文本")
    return text


def clean_text(text: str) -> str:
    """基础清洗：去空行/空白符归一化"""
    lines = [line.strip() for line in text.splitlines()]
    # 去掉空行与纯符号行，保留段落结构
    cleaned = [line for line in lines if line and not set(line) <= set("-\t *_#|·•")]
    return "\n".join(cleaned)
